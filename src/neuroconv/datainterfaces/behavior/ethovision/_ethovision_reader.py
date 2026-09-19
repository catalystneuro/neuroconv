"""Read Noldus EthoVision XT Track exports across Excel, CSV, and TXT containers."""

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np

TRACK_SHEET_PATTERN = re.compile(r"^Track-(?P<arena>.+)-Subject (?P<subject>.+)$")
SUPPORTED_SUFFIXES = (".xlsx", ".csv", ".txt")

TRIAL_TIME_COLUMN = "Trial time"
RECORDING_TIME_COLUMN = "Recording time"


@dataclass(frozen=True)
class EthoVisionTrackSource:
    """The identity and container-local name of one EthoVision Track."""

    arena: str
    subject: str
    source_name: str


@dataclass
class EthoVisionTrackData:
    """One Track's header fields, time vectors, units, and per-frame channels."""

    header: dict[str, str | None]
    units: dict[str, str | None]
    trial_time: np.ndarray
    recording_time: np.ndarray
    channels: dict[str, np.ndarray] = field(default_factory=dict)


@dataclass
class EthoVisionScoringEvent:
    """One manual-scoring point event or paired state bout."""

    subject: str
    behavior: str
    onset: float
    duration: float | None


def get_available_tracks(file_path) -> list[dict[str, str]]:
    """Return the complete selector arguments for every available Track."""
    return [
        {"arena_name": source.arena, "subject_name": source.subject}
        for source in _get_track_sources(file_path=file_path)
    ]


def select_track_source(
    file_path,
    *,
    arena_name: str | None = None,
    subject_name: str | None = None,
) -> EthoVisionTrackSource:
    """Resolve exactly one Track, inferring selectors only when the result is unambiguous."""
    available = _get_track_sources(file_path=file_path)
    selected = [
        source
        for source in available
        if (arena_name is None or source.arena == arena_name)
        and (subject_name is None or source.subject == subject_name)
    ]
    if len(selected) == 1:
        return selected[0]

    available_identities = [(source.arena, source.subject) for source in available]
    if not selected:
        raise ValueError(
            f"No EthoVision Track matches arena_name={arena_name!r}, subject_name={subject_name!r} "
            f"in '{file_path}'. Available tracks: {available_identities}."
        )
    raise ValueError(
        f"arena_name={arena_name!r}, subject_name={subject_name!r} does not identify one Track in "
        f"'{file_path}'. Matching tracks: {[(source.arena, source.subject) for source in selected]}."
    )


def read_track(file_path, *, source: EthoVisionTrackSource) -> EthoVisionTrackData:
    """Read one selected Track using the loader for its source container."""
    path = _validate_file_path(file_path=file_path)
    if path.suffix.lower() == ".xlsx":
        rows = _read_excel_rows(file_path=path, sheet_name=source.source_name)
    else:
        rows = _read_delimited_rows(file_path=path)
    return _track_data_from_rows(rows=rows, source_name=source.source_name)


def read_scoring_events(file_path, *, arena_name: str) -> list[EthoVisionScoringEvent]:
    """Read an arena's Manual Scoring sheet when the source is an Excel workbook."""
    path = _validate_file_path(file_path=file_path)
    if path.suffix.lower() != ".xlsx":
        return []

    sheet_name = f"Manual Scoring-{arena_name}"
    workbook = _load_workbook(file_path=path)
    try:
        if sheet_name not in workbook.sheetnames:
            return []
        rows = [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
    finally:
        workbook.close()
    return _scoring_events_from_rows(rows=rows, source_name=sheet_name)


def _get_track_sources(file_path) -> list[EthoVisionTrackSource]:
    path = _validate_file_path(file_path=file_path)
    if path.suffix.lower() == ".xlsx":
        workbook = _load_workbook(file_path=path)
        try:
            sources = []
            for sheet_name in workbook.sheetnames:
                match = TRACK_SHEET_PATTERN.match(sheet_name)
                if match is not None:
                    sources.append(
                        EthoVisionTrackSource(
                            arena=match.group("arena"),
                            subject=f"Subject {match.group('subject')}",
                            source_name=sheet_name,
                        )
                    )
            return sources
        finally:
            workbook.close()

    rows = _read_delimited_rows(file_path=path)
    header, column_names, _units, _data_rows = _split_header_and_table(rows=rows, source_name=path.name)
    _validate_track_columns(column_names=column_names, source_name=path.name)
    arena = header.get("Arena name")
    subject = header.get("Subject name")
    if not arena or not subject:
        raise ValueError(f"'{path}' does not declare both 'Arena name' and 'Subject name' in its header.")
    return [EthoVisionTrackSource(arena=str(arena), subject=str(subject), source_name=path.name)]


def _track_data_from_rows(rows: list[list], *, source_name: str) -> EthoVisionTrackData:
    header, column_names, units, data_rows = _split_header_and_table(rows=rows, source_name=source_name)
    _validate_track_columns(column_names=column_names, source_name=source_name)
    columns = {
        name: np.asarray(
            [_parse_track_value(row[index] if index < len(row) else None) for row in data_rows],
            dtype=float,
        )
        for index, name in enumerate(column_names)
    }
    trial_time = columns.pop(TRIAL_TIME_COLUMN)
    recording_time = columns.pop(RECORDING_TIME_COLUMN)
    return EthoVisionTrackData(
        header=header,
        units={name: _normalize_unit(units.get(name)) for name in columns},
        trial_time=trial_time,
        recording_time=recording_time,
        channels=columns,
    )


def _scoring_events_from_rows(rows: list[list], *, source_name: str) -> list[EthoVisionScoringEvent]:
    _header, column_names, _units, data_rows = _split_header_and_table(rows=rows, source_name=source_name)
    required_columns = {RECORDING_TIME_COLUMN, "Subject", "Behavior", "Event"}
    missing_columns = required_columns.difference(column_names)
    if missing_columns:
        raise ValueError(f"'{source_name}' is missing Manual Scoring columns: {sorted(missing_columns)}.")

    subject_index = column_names.index("Subject")
    behavior_index = column_names.index("Behavior")
    event_index = column_names.index("Event")
    recording_time_index = column_names.index(RECORDING_TIME_COLUMN)

    events: list[EthoVisionScoringEvent] = []
    open_bouts: dict[tuple[str, str], float] = {}
    for row in data_rows:
        subject, behavior, event = row[subject_index], row[behavior_index], row[event_index]
        onset = float(row[recording_time_index])
        key = (subject, behavior)
        if event == "point event":
            events.append(EthoVisionScoringEvent(subject=subject, behavior=behavior, onset=onset, duration=None))
        elif event == "state start":
            open_bouts[key] = onset
        elif event == "state stop":
            if key not in open_bouts:
                raise ValueError(
                    f"'{source_name}' has a 'state stop' for subject '{subject}', behavior '{behavior}' "
                    f"at {RECORDING_TIME_COLUMN}={onset} with no preceding 'state start' to close."
                )
            start = open_bouts.pop(key)
            events.append(
                EthoVisionScoringEvent(subject=subject, behavior=behavior, onset=start, duration=onset - start)
            )
        else:
            raise ValueError(f"'{source_name}' has an unrecognized Event value '{event}'.")

    for (subject, behavior), onset in open_bouts.items():
        events.append(EthoVisionScoringEvent(subject=subject, behavior=behavior, onset=onset, duration=np.nan))
    return events


def _validate_file_path(file_path) -> Path:
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported EthoVision suffix '{path.suffix}'. Expected one of {SUPPORTED_SUFFIXES}.")
    return path


def _load_workbook(file_path: Path):
    import openpyxl

    return openpyxl.load_workbook(file_path, read_only=True, data_only=True)


def _read_excel_rows(file_path: Path, *, sheet_name: str) -> list[list]:
    workbook = _load_workbook(file_path=file_path)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"'{sheet_name}' is not a sheet in this file. Found: {workbook.sheetnames}")
        return [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
    finally:
        workbook.close()


def _read_delimited_rows(file_path: Path) -> list[list[str | None]]:
    raw = file_path.read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        text = raw.decode("utf-16")
    else:
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1252")

    delimiter = "," if file_path.suffix.lower() == ".csv" else _detect_delimiter(text=text, file_path=file_path)
    rows = []
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        rows.append([value if value != "" else None for value in row])
    return rows


def _detect_delimiter(*, text: str, file_path: Path) -> str:
    sample = "\n".join(text.splitlines()[:50])
    try:
        return csv.Sniffer().sniff(sample, delimiters=";\t,").delimiter
    except csv.Error as exception:
        raise ValueError(f"Could not determine the delimiter used by '{file_path}'.") from exception


def _split_header_and_table(rows: list[list], *, source_name: str):
    """Split self-declaring EthoVision rows into header, names, units, and data."""
    if not rows or len(rows[0]) < 2 or rows[0][0] != "Number of header lines:":
        raise ValueError(f"'{source_name}' does not begin with an EthoVision 'Number of header lines:' row.")
    try:
        header_lines = int(rows[0][1])
    except (TypeError, ValueError) as exception:
        raise ValueError(f"'{source_name}' has an invalid Number of header lines value: {rows[0][1]!r}.") from exception
    if len(rows) < header_lines:
        raise ValueError(f"'{source_name}' declares {header_lines} header lines but contains only {len(rows)} rows.")

    header = {
        row[0]: (row[1] if len(row) > 1 else None) for row in rows[1 : header_lines - 2] if row and row[0] is not None
    }
    column_names = [name for name in rows[header_lines - 2] if name is not None]
    units_row = rows[header_lines - 1]
    units = {name: units_row[index] if index < len(units_row) else None for index, name in enumerate(column_names)}
    return header, column_names, units, rows[header_lines:]


def _validate_track_columns(*, column_names: list[str], source_name: str) -> None:
    missing_columns = {TRIAL_TIME_COLUMN, RECORDING_TIME_COLUMN}.difference(column_names)
    if missing_columns:
        raise ValueError(f"'{source_name}' is not a Track export; missing columns: {sorted(missing_columns)}.")


def _parse_track_value(value) -> float:
    """Parse a numeric Track value, preserving missing samples as ``numpy.nan``."""
    if value is None or value in {"-", "NA"}:
        return np.nan
    return float(value)


def _normalize_unit(unit: str | None) -> str | None:
    """Normalize the optional parentheses used around units in delimited exports."""
    if unit is None:
        return None
    return unit[1:-1] if unit.startswith("(") and unit.endswith(")") else unit
